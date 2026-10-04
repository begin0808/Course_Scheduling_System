"""跨節次表教師的「顯示」:國中小(國中 45 分、國小 40 分)同一位老師兩部都有課。

衝堂判定本來就以牆鐘時間為準(D7,見 test_timetables);這裡驗的是使用者看到的東西——
教師課表與通知單的格線只能照一套節次表畫,另一套的課必須標出實際時間,
衝堂訊息也要寫出對方那一節真正的時間。
"""

import io

import pytest
from openpyxl import load_workbook

from app.models.user import Role
from tests.conftest import make_user
from tests.dates import SEM_END, SEM_START, WED
from tests.test_substitutions import _World

PW = "password123"


@pytest.fixture
def k9(env):
    """國中小:預設國中節次表 + 國小節次表;林師在 701(國中)與 601(國小)都有英語。"""
    client, db = env
    make_user(db, "s", PW, roles=[Role.scheduler])
    client.post("/api/auth/login", json={"username": "s", "password": PW})
    sid = client.post("/api/semesters", json={
        "academic_year": 115, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()["id"]
    w = _World(client, db, sid)
    table = client.post(f"/api/semesters/{sid}/period-tables",
                        json={"name": "國小節次表", "template_key": "elementary"}).json()
    w.classes["601"] = client.post(f"/api/class-units{w.q}", json={
        "grade": 6, "name": "601", "track": "elementary", "period_table_id": table["id"],
    }).json()["id"]
    w.teacher("林師", ["英語"])
    w.teacher("陳師", ["英語"])
    w.place("林師", "英語", "701", 0)   # 國中第一節 08:20–09:05
    w.place("林師", "英語", "601", 3)   # 國小第四節 11:10–11:50(國中第四節是 11:15–12:00)
    return w


def test_conflict_message_gives_the_other_lessons_real_time(k9):
    """國小第一節 08:30 撞到國中第一節 08:20–09:05:訊息要寫出那堂課實際的時間。"""
    w = k9
    slots = [p for p in w.client.get(
        f"/api/class-units/{w.classes['601']}/period-table").json()["periods"]
        if p["weekday"] == 3 and p["type"] == "regular"]
    a = w.client.post(f"/api/assignments{w.q}", json={
        "class_id": w.classes["601"], "subject_id": w.subject("英語"), "periods_per_week": 1,
        "teachers": [{"teacher_id": w.teachers["林師"]}], "block_rules": [],
    }).json()
    r = w.client.post(f"/api/timetables/{w.tt}/entries", json={
        "course_assignment_id": a["id"], "weekday": 3,
        "period_no": slots[0]["period_no"], "span": 1})
    assert r.status_code == 409
    messages = [c["message"] for c in r.json()["detail"]["conflicts"]]
    assert any("701 班英語(08:20–09:05)" in m for m in messages), messages


def test_teacher_timetable_marks_the_real_time_of_other_table_lessons(k9):
    w = k9
    w.publish()
    r = w.client.get(
        f"/api/export/timetable{w.q}&view=teacher&target_id={w.teachers['林師']}&fmt=xlsx")
    assert r.status_code == 200, r.text
    cells = [str(v) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True) for v in row if v]
    c601 = next(c for c in cells if "601" in c)
    c701 = next(c for c in cells if "701" in c)
    assert "11:10–11:50" in c601          # 國小的課:標出實際時間
    assert "–" not in c701                # 國中的課就在自己的格線上,不必標


def test_class_timetable_needs_no_note(k9):
    """班級課表照該班自己的節次表畫,時間本來就對。"""
    w = k9
    w.publish()
    r = w.client.get(
        f"/api/export/timetable{w.q}&view=class&target_id={w.classes['601']}&fmt=xlsx")
    cells = [str(v) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True) for v in row if v]
    assert "–" not in next(c for c in cells if "林師" in c)


def test_substitute_slip_marks_the_real_time_of_other_table_lessons(k9):
    """陳師同一張代課單上有國中與國小的課:格線照第一節(國中)畫,國小那格要標時間。"""
    w = k9
    w.publish()
    affected = w.leave("林師")
    for ap in affected:
        status, _ = w.assign(ap["id"], type="substitute", handler_teacher_id=w.teachers["陳師"])
        assert status == 200
    q = "&".join(f"affected_period_ids={ap['id']}" for ap in affected)
    body = w.client.get(f"/api/substitute-slips{w.q}&{q}").json()

    teacher = next(s for s in body["slips"] if s["kind"] == "teacher")
    notes = {c["actor"]: c["time_note"] for c in teacher["weeks"][0]["cells"]}
    assert notes == {"701[代]": "", "601[代]": "11:10–11:50"}
    assert WED.isoformat() == teacher["weeks"][0]["cells"][0]["date"]

    for slip in (s for s in body["slips"] if s["kind"] == "class"):
        assert all(c["time_note"] == "" for c in slip["weeks"][0]["cells"])   # 班級單照自己的格線


# ── 完全中學:兩部節數不同(國中七節、高中八節)──────────────────
@pytest.fixture
def k12(env):
    """預設國中節次表(七節)+ 高中部節次表(八節);林師在 901 與高一甲都有英語。"""
    client, db = env
    make_user(db, "s", PW, roles=[Role.scheduler])
    client.post("/api/auth/login", json={"username": "s", "password": PW})
    sid = client.post("/api/semesters", json={
        "academic_year": 115, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()["id"]
    w = _World(client, db, sid)
    table = client.post(f"/api/semesters/{sid}/period-tables",
                        json={"name": "高中部節次表", "template_key": "senior_high"}).json()
    w.classes["高一甲"] = client.post(f"/api/class-units{w.q}", json={
        "grade": 10, "name": "高一甲", "track": "senior_high", "period_table_id": table["id"],
    }).json()["id"]
    w.teacher("林師", ["英語"])
    w.teacher("陳師", ["英語"])
    w.place("林師", "英語", "901", 0)      # 國中第一節
    w.place("林師", "英語", "高一甲", 7)   # 高中第八節 16:20–17:10:國中節次表沒有這一列
    w.publish()
    return w


def test_teacher_timetable_adds_the_row_the_default_table_lacks(k12):
    """高中第八節的課不能因為國中格線只有七列就從教師課表上消失。"""
    w = k12
    r = w.client.get(
        f"/api/export/timetable{w.q}&view=teacher&target_id={w.teachers['林師']}&fmt=xlsx")
    assert r.status_code == 200, r.text
    cells = [str(v) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True) for v in row if v]
    assert "第八節" in cells
    eighth = next(c for c in cells if "高一甲" in c)
    assert "16:20–17:10" in eighth


def test_teacher_without_cross_table_lessons_gets_no_extra_row(k12):
    """只在國中上課的老師,課表維持七列,不會多一列空的第八節。"""
    w = k12
    w2_teacher = w.teachers["陳師"]
    r = w.client.get(f"/api/export/timetable{w.q}&view=teacher&target_id={w2_teacher}&fmt=xlsx")
    cells = [str(v) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True) for v in row if v]
    assert "第七節" in cells and "第八節" not in cells


def test_substitute_slip_uses_the_table_with_more_rows(k12):
    """教師代課單的格線照節數多的那套(高中八節)畫,第八節那格才有地方放。"""
    w = k12
    affected = w.leave("林師")
    for ap in affected:
        status, _ = w.assign(ap["id"], type="substitute", handler_teacher_id=w.teachers["陳師"])
        assert status == 200
    q = "&".join(f"affected_period_ids={ap['id']}" for ap in affected)
    body = w.client.get(f"/api/substitute-slips{w.q}&{q}").json()

    teacher = next(s for s in body["slips"] if s["kind"] == "teacher")
    assert len(teacher["rows"]) == 8
    cells = {c["actor"]: (c["ordinal"], c["time_note"]) for c in teacher["weeks"][0]["cells"]}
    # 格線是高中的:高中那格不必標時間,國中那格標實際時間
    assert cells == {"901[代]": (1, "08:20–09:05"), "高一甲[代]": (8, "")}


# ── 第八節:有的學校有、有的沒有,國中/高中各種組合 ──────────────
def _add_eighth_period(w, table, start, end):
    """在節次表最後加一列第八節(等同「編輯節次表 → 新增節次列」)。時間可以不填。"""
    periods = [{k: p[k] for k in ("weekday", "period_no", "name", "start_time", "end_time", "type")}
               for p in table["periods"]]
    for wd in range(1, 6):
        periods.append({"weekday": wd, "period_no": 10, "name": "第八節",
                        "start_time": start, "end_time": end, "type": "regular"})
    r = w.client.put(f"/api/period-tables/{table['id']}/periods", json=periods)
    assert r.status_code == 200, r.text


def _k12_with_junior_eighth(env, start, end):
    client, db = env
    make_user(db, "s", PW, roles=[Role.scheduler])
    client.post("/api/auth/login", json={"username": "s", "password": PW})
    sem = client.post("/api/semesters", json={
        "academic_year": 115, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()
    w = _World(client, db, sem["id"])
    _add_eighth_period(w, sem["period_tables"][0], start, end)
    table = client.post(f"/api/semesters/{sem['id']}/period-tables",
                        json={"name": "高中部節次表", "template_key": "senior_high"}).json()
    w.classes["高一甲"] = client.post(f"/api/class-units{w.q}", json={
        "grade": 10, "name": "高一甲", "track": "senior_high", "period_table_id": table["id"],
    }).json()["id"]
    w.teacher("林師", ["英語"])
    return w


def _try_place(w, klass, period_idx, weekday=3):
    slots = [p for p in w.client.get(
        f"/api/class-units/{w.klass(klass)}/period-table").json()["periods"]
        if p["weekday"] == weekday and p["type"] == "regular"]
    a = w.client.post(f"/api/assignments{w.q}", json={
        "class_id": w.klass(klass), "subject_id": w.subject("英語"), "periods_per_week": 1,
        "teachers": [{"teacher_id": w.teachers["林師"]}], "block_rules": [],
    }).json()
    return w.client.post(f"/api/timetables/{w.tt}/entries", json={
        "course_assignment_id": a["id"], "weekday": weekday,
        "period_no": slots[period_idx]["period_no"], "span": 1})


def test_both_divisions_have_an_eighth_period(env):
    """國中、高中都有第八節(16:05–16:50 與 16:20–17:10):同一天會擋,課表只有一列第八節。"""
    w = _k12_with_junior_eighth(env, "16:05:00", "16:50:00")
    assert _try_place(w, "901", 7).status_code == 201
    r = _try_place(w, "高一甲", 7)
    assert r.status_code == 409
    assert "901 班英語(16:05–16:50)" in r.json()["detail"]["conflicts"][0]["message"]
    assert _try_place(w, "高一甲", 7, weekday=4).status_code == 201   # 換一天就可以

    w.publish()
    r = w.client.get(
        f"/api/export/timetable{w.q}&view=teacher&target_id={w.teachers['林師']}&fmt=xlsx")
    cells = [str(v) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True) for v in row if v]
    assert cells.count("第八節") == 1
    assert "16:20–17:10" in next(c for c in cells if "高一甲" in c)
    assert "–" not in next(c for c in cells if "901" in c)


def test_period_without_times_still_blocks_the_same_period_number(env):
    """學校自己加的第八節沒填起訖時間:無從比時間,退回節次號——不能讓老師同時段兩堂課。"""
    w = _k12_with_junior_eighth(env, None, None)
    assert _try_place(w, "901", 7).status_code == 201
    r = _try_place(w, "高一甲", 7)
    assert r.status_code == 409, r.json()
    assert _try_place(w, "高一甲", 6).status_code == 201   # 不同節次號照常可排
