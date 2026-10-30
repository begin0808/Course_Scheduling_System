"""代導(v1.4.0,LINE 社群使用者需求):導師請假那天,由誰代他帶班。

先前請假只展開「有課的節次」;導師當天沒課就什麼都不會出現,班上卻照樣需要有人帶。
使用者定案的規則:以天為單位、半天也要代導、全天/半天分開統計(代導費由學校換算)、
要代導通知單、誰都能代導(有課也可以,只排除整天請假的人)。
"""

import io
from datetime import time, timedelta

from openpyxl import load_workbook
from sqlalchemy import select

from app.core import clock
from app.models.notification import Notification
from tests.dates import FRI, MON, SAT, SUN, THU, TUE, WED
from tests.test_substitutions import _World, env2  # noqa: F401 - env2 為 fixture


def _homeroom(w: _World, klass: str, teacher: str, grade: int = 7) -> None:
    if klass not in w.classes:
        w.classes[klass] = w.client.post(f"/api/class-units{w.q}", json={
            "grade": grade, "name": klass, "track": "junior_high"}).json()["id"]
    r = w.client.patch(f"/api/class-units/{w.classes[klass]}", json={
        "grade": grade, "name": klass, "track": "junior_high",
        "homeroom_teacher_id": w.teachers[teacher]})
    assert r.status_code == 200, r.json()


def _leave(w: _World, teacher: str, start, end=None, *, start_time=None, end_time=None,
           leave_type: str = "personal") -> dict:
    r = w.client.post(f"/api/leaves{w.q}", json={
        "teacher_id": w.teachers[teacher], "leave_type": leave_type,
        "start_date": start.isoformat(), "end_date": (end or start).isoformat(),
        "start_time": start_time, "end_time": end_time})
    assert r.status_code == 201, r.json()
    return r.json()


def _covers(w: _World, leave_id: int) -> list[dict]:
    return w.client.get(f"/api/leaves/{leave_id}").json()["homeroom_covers"]


def _assign(w: _World, cover_id: int, **body):
    r = w.client.put(f"/api/homeroom-covers/{cover_id}", json=body)
    return r.status_code, r.json()


def _notes(w: _World, teacher: str) -> list[str]:
    return [n.title for n in w.db.scalars(
        select(Notification).where(Notification.teacher_id == w.teachers[teacher])
        .order_by(Notification.id))]


# ── 展開 ──────────────────────────────────────────
def test_homeroom_leave_creates_one_cover_per_school_day(env2):  # noqa: F811
    """週一到週日的假:只有週一到週五要代導;課表沒發布、當天沒課也照樣展開。"""
    w = env2
    w.teacher("王師")
    _homeroom(w, "701", "王師")
    leave = _leave(w, "王師", MON, SUN)

    assert leave["affected_count"] == 0          # 王師一節課都沒有
    covers = leave["homeroom_covers"]
    assert [c["date"] for c in covers] == [d.isoformat() for d in (MON, TUE, WED, THU, FRI)]
    assert {c["class_name"] for c in covers} == {"701"}
    assert {c["part_label"] for c in covers} == {"全天"}
    assert {c["status"] for c in covers} == {"pending"}
    assert leave["cover_pending_count"] == 5


def test_non_homeroom_teacher_has_no_cover(env2):  # noqa: F811
    w = env2
    w.teacher("陳師")
    assert _leave(w, "陳師", WED)["homeroom_covers"] == []


def test_teacher_with_two_homeroom_classes_gets_both(env2):  # noqa: F811
    w = env2
    w.teacher("王師")
    _homeroom(w, "701", "王師")
    _homeroom(w, "702", "王師")
    covers = _leave(w, "王師", WED)["homeroom_covers"]
    assert [c["class_name"] for c in covers] == ["701", "702"]


def test_half_day_leave_is_marked_morning_or_afternoon(env2):  # noqa: F811
    """半天假也要代導;跨日的假,第一天下午、中間整天、最後一天上午。"""
    w = env2
    w.teacher("王師")
    _homeroom(w, "701", "王師")
    morning = _leave(w, "王師", MON, end_time="12:00")["homeroom_covers"]
    assert [(c["part_label"], c["is_half_day"]) for c in morning] == [("上午", True)]

    span = _leave(w, "王師", TUE, THU, start_time="13:00", end_time="12:00")["homeroom_covers"]
    assert [c["part_label"] for c in span] == ["下午", "全天", "上午"]
    assert [c["is_half_day"] for c in span] == [True, False, True]


# ── 人選 ──────────────────────────────────────────
def _school(w: _World) -> int:
    """王師(701 導師)週三請整天假。陳師 702 導師(同年級)、李師 801 導師、林師教 701 數學
    (而且週三第一節正在上課)、張師沒有任何關係、趙師週三整天請假、吳師週三下午請假。"""
    for name in ("王師", "陳師", "李師", "林師", "張師", "趙師", "吳師"):
        w.teacher(name, ["數學"] if name == "林師" else [])
    _homeroom(w, "701", "王師")
    _homeroom(w, "702", "陳師")
    _homeroom(w, "801", "李師", grade=8)
    w.place("林師", "數學", "701", 0)
    w.publish()
    _leave(w, "趙師", WED)
    _leave(w, "吳師", WED, start_time="13:00")
    return _leave(w, "王師", WED)["homeroom_covers"][0]["id"]


def test_options_same_grade_homeroom_then_class_teacher_then_others(env2):  # noqa: F811
    w = env2
    cover_id = _school(w)
    r = w.client.get(f"/api/homeroom-covers/{cover_id}/options")
    assert r.status_code == 200, r.json()
    got = [(o["teacher_name"], o["group_label"], o["note"]) for o in r.json()["options"]]
    assert got[:2] == [("陳師", "同年級導師", "702 導師"), ("林師", "該班任課老師", "教 701 數學")]
    others = [name for name, group, _ in got[2:]]
    assert {group for _, group, _ in got[2:]} == {"其他老師"}
    # 李師是 8 年級導師 → 其他老師;吳師只請半天,另外半天可以代導
    assert sorted(others) == sorted(["李師", "張師", "吳師"])
    names = [g[0] for g in got]
    assert "王師" not in names and "趙師" not in names   # 請假本人、整天請假的人不列


def test_teacher_who_has_class_can_still_cover(env2):  # noqa: F811
    """代導不是上課:林師週三第一節有課,照樣可以代導。"""
    w = env2
    cover_id = _school(w)
    code, body = _assign(w, cover_id, handler_teacher_id=w.teachers["林師"])
    assert code == 200, body
    cover = _covers(w, w.client.get(f"/api/leaves{w.q}").json()[0]["id"])[0]
    assert (cover["status"], cover["handler_name"]) == ("resolved", "林師")
    assert "代導通知" in _notes(w, "林師")[-1]


def test_cannot_assign_teacher_on_full_day_leave_or_the_absent_one(env2):  # noqa: F811
    w = env2
    cover_id = _school(w)
    code, body = _assign(w, cover_id, handler_teacher_id=w.teachers["趙師"])
    assert code == 409 and "整天請假" in body["detail"]
    code, body = _assign(w, cover_id, handler_teacher_id=w.teachers["王師"])
    assert code == 409 and "自己的班" in body["detail"]
    code, body = _assign(w, cover_id)
    assert code == 409 and "請選擇代導老師" in body["detail"]


# ── 指派 ──────────────────────────────────────────
def test_whole_leave_assigns_every_pending_day_and_skips_days_off(env2):  # noqa: F811
    """多天的假一次派完;代導老師整天請假的那天跳過、仍待代導。"""
    w = env2
    w.teacher("王師")
    w.teacher("陳師")
    _homeroom(w, "701", "王師")
    _leave(w, "陳師", THU)
    leave = _leave(w, "王師", MON, FRI)
    covers = leave["homeroom_covers"]

    code, body = _assign(w, covers[0]["id"], handler_teacher_id=w.teachers["陳師"],
                         whole_leave=True)
    assert code == 200, body
    assert body["assigned"] == 4
    assert [(s["date"], "整天請假" in s["reason"]) for s in body["skipped"]] == [
        (THU.isoformat(), True)]
    after = _covers(w, leave["id"])
    assert [c["status"] for c in after] == ["resolved", "resolved", "resolved", "pending",
                                           "resolved"]
    # 代導老師只收到每天各一則通知,不會重複
    assert len([t for t in _notes(w, "陳師") if "代導通知" in t]) == 4


def test_no_cover_and_undo(env2):  # noqa: F811
    w = env2
    w.teacher("王師")
    w.teacher("陳師")
    _homeroom(w, "701", "王師")
    leave = _leave(w, "王師", WED)
    cover_id = leave["homeroom_covers"][0]["id"]

    code, _ = _assign(w, cover_id, no_cover=True)
    assert code == 200
    cover = _covers(w, leave["id"])[0]
    assert (cover["status"], cover["no_cover"], cover["handler_name"]) == ("resolved", True, None)

    _assign(w, cover_id, handler_teacher_id=w.teachers["陳師"])
    assert w.client.delete(f"/api/homeroom-covers/{cover_id}").status_code == 200
    cover = _covers(w, leave["id"])[0]
    assert (cover["status"], cover["no_cover"], cover["handler_name"]) == ("pending", False, None)
    assert _notes(w, "陳師")[-1] == "原訂代導已取消"


def test_reassign_tells_the_previous_teacher(env2):  # noqa: F811
    w = env2
    for name in ("王師", "陳師", "林師"):
        w.teacher(name)
    _homeroom(w, "701", "王師")
    cover_id = _leave(w, "王師", WED)["homeroom_covers"][0]["id"]
    _assign(w, cover_id, handler_teacher_id=w.teachers["陳師"])
    _assign(w, cover_id, handler_teacher_id=w.teachers["林師"])
    assert _notes(w, "陳師")[-1] == "原訂代導已取消"
    assert "代導通知" in _notes(w, "林師")[-1]


# ── 銷假 ──────────────────────────────────────────
def test_cancel_leave_cancels_future_covers_and_keeps_days_already_covered(env2, monkeypatch):  # noqa: F811
    w = env2
    w.teacher("王師")
    w.teacher("陳師")
    _homeroom(w, "701", "王師")
    leave = _leave(w, "王師", MON, WED)
    covers = leave["homeroom_covers"]
    _assign(w, covers[0]["id"], handler_teacher_id=w.teachers["陳師"], whole_leave=True)

    # 週二上班時銷假:週一已經代過了(保留、照算),週二、週三取消
    monkeypatch.setattr(clock, "school_today", lambda: TUE)
    r = w.client.post(f"/api/leaves/{leave['id']}/cancel")
    assert r.status_code == 200, r.json()
    assert r.json()["revoked_cover_count"] == 2
    assert "陳師" in r.json()["notified_teachers"]
    statuses = [c["status"] for c in _covers(w, leave["id"])]
    assert statuses == ["completed", "cancelled", "cancelled"]
    assert _notes(w, "陳師")[-1].startswith("原訂代導已取消")


def test_past_day_cannot_be_changed(env2, monkeypatch):  # noqa: F811
    """已安排、那天已經過了:不能改派。(還沒安排的過去日子可以補登,見 test_backfill)"""
    w = env2
    w.teacher("王師")
    w.teacher("陳師")
    w.teacher("林師")
    _homeroom(w, "701", "王師")
    cover_id = _leave(w, "王師", MON)["homeroom_covers"][0]["id"]
    _assign(w, cover_id, handler_teacher_id=w.teachers["陳師"])
    monkeypatch.setattr(clock, "school_today", lambda: TUE)
    code, body = _assign(w, cover_id, handler_teacher_id=w.teachers["林師"])
    assert code == 409 and "已經過了" in body["detail"]


# ── 月結、通知單、看板 ──────────────────────────────
def test_monthly_report_counts_full_and_half_days_separately(env2):  # noqa: F811
    """同一份月結報表:陳師代導 2 天全天 + 1 個半天;只代導、沒代課的人也要列出來。"""
    w = env2
    for name in ("王師", "林師", "陳師"):
        w.teacher(name)
    _homeroom(w, "701", "王師")
    _homeroom(w, "702", "林師")
    full = _leave(w, "王師", MON, TUE)["homeroom_covers"]
    half = _leave(w, "林師", WED, end_time="12:00")["homeroom_covers"]
    _assign(w, full[0]["id"], handler_teacher_id=w.teachers["陳師"], whole_leave=True)
    _assign(w, half[0]["id"], handler_teacher_id=w.teachers["陳師"])

    # 基準週一到週三必在同一個月(tests.dates.base_monday)
    report = w.client.get(
        f"/api/substitution-stats{w.q}&year={MON.year}&month={MON.month}").json()
    chen = next(s for s in report["summaries"] if s["teacher_name"] == "陳師")
    assert (chen["cover_full_days"], chen["cover_half_days"], chen["handled_count"]) == (2, 1, 0)
    assert [(d["class_name"], d["part_label"]) for d in report["cover_details"]] == [
        ("701", "全天"), ("701", "全天"), ("702", "上午")]

    x = w.client.get(f"/api/substitution-stats/export{w.q}&year={MON.year}&month={MON.month}")
    wb = load_workbook(io.BytesIO(x.content))
    summary = [[c.value for c in row] for row in wb["彙總"].iter_rows()]
    assert summary[0][5:7] == ["代導全天", "代導半天"]
    assert summary[1][0] == "陳師" and summary[1][5:7] == [2, 1]
    covers = [[c.value for c in row] for row in wb["代導明細"].iter_rows()]
    assert covers[0] == ["教師", "日期", "班級", "時段", "原導師", "假別"]
    assert covers[3][2:5] == ["702", "上午", "林師"]


def test_no_cover_days_are_not_counted(env2):  # noqa: F811
    w = env2
    w.teacher("王師")
    _homeroom(w, "701", "王師")
    cover_id = _leave(w, "王師", WED)["homeroom_covers"][0]["id"]
    _assign(w, cover_id, no_cover=True)
    r = w.client.get(f"/api/substitution-stats{w.q}&year={WED.year}&month={WED.month}").json()
    assert r["summaries"] == [] and r["cover_details"] == []


def test_slips_one_per_cover_teacher(env2):  # noqa: F811
    w = env2
    for name in ("王師", "陳師", "林師"):
        w.teacher(name)
    _homeroom(w, "701", "王師")
    covers = _leave(w, "王師", MON, WED, leave_type="sick")["homeroom_covers"]
    _assign(w, covers[0]["id"], handler_teacher_id=w.teachers["陳師"])
    _assign(w, covers[1]["id"], handler_teacher_id=w.teachers["林師"])
    _assign(w, covers[2]["id"], handler_teacher_id=w.teachers["陳師"])
    ids = "&".join(f"cover_ids={c['id']}" for c in covers)
    r = w.client.get(f"/api/homeroom-cover-slips{w.q}&{ids}")
    assert r.status_code == 200, r.json()
    slips = r.json()["slips"]
    assert [s["teacher_name"] for s in slips] == sorted(["陳師", "林師"])
    chen = next(s for s in slips if s["teacher_name"] == "陳師")
    assert [(row["date"], row["class_name"], row["absent_teacher_name"], row["leave_type_label"])
            for row in chen["rows"]] == [
        (MON.isoformat(), "701", "王師", "病假"), (WED.isoformat(), "701", "王師", "病假")]


def test_daily_board_and_log_include_covers(env2):  # noqa: F811
    w = env2
    w.teacher("王師")
    w.teacher("陳師")
    _homeroom(w, "701", "王師")
    covers = _leave(w, "王師", WED, THU)["homeroom_covers"]
    _assign(w, covers[0]["id"], handler_teacher_id=w.teachers["陳師"])

    board = w.client.get(f"/api/daily-board{w.q}&on={WED.isoformat()}").json()
    assert [(c["class_name"], c["handler_name"], c["absent_teacher_name"])
            for c in board["homeroom_covers"]] == [("701", "陳師", "王師")]

    log = w.client.get(f"/api/homeroom-covers{w.q}&teacher_id={w.teachers['陳師']}").json()
    assert [c["date"] for c in log] == [WED.isoformat()]
    everyone = w.client.get(
        f"/api/homeroom-covers{w.q}&date_from={WED.isoformat()}"
        f"&date_to={(WED + timedelta(days=1)).isoformat()}").json()
    assert len(everyone) == 2


def test_saturday_counts_only_for_six_day_schools(env2):  # noqa: F811
    """上課日由節次表決定:預設五日制,週六不必代導。"""
    w = env2
    w.teacher("王師")
    _homeroom(w, "701", "王師")
    covers = _leave(w, "王師", SAT, SUN)["homeroom_covers"]
    assert covers == []


def test_leave_time_window_values_are_kept(env2):  # noqa: F811
    """直接看資料表:半天假記下時段,全天的兩端都是空的。"""
    from app.models.leave import HomeroomCover

    w = env2
    w.teacher("王師")
    _homeroom(w, "701", "王師")
    _leave(w, "王師", WED, start_time="10:00", end_time="14:00")
    cover = w.db.scalar(select(HomeroomCover))
    assert (cover.from_time, cover.to_time) == (time(10, 0), time(14, 0))
    assert cover.part_label == "10:00~14:00"
