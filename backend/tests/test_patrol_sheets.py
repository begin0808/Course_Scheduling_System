"""巡堂表(v1.2.12):已發布課表疊上當天的調代課,印出那一天每班每節真正的樣子。"""

import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from app.models.user import Role
from tests.conftest import make_user
from tests.dates import SEM_END, SEM_START, WED
from tests.test_substitutions import _swap_first, _World

PW = "password123"


@pytest.fixture
def w(env):
    client, db = env
    make_user(db, "s", PW, roles=[Role.scheduler])
    client.post("/api/auth/login", json={"username": "s", "password": PW})
    sid = client.post("/api/semesters", json={
        "academic_year": 115, "term": 1, "template_key": "junior_high",
        "start_date": SEM_START.isoformat(), "end_date": SEM_END.isoformat(),
    }).json()["id"]
    return _World(client, db, sid)


def _sheets(w, day=WED, to=None):
    q = f"{w.q}&date_from={day.isoformat()}" + (f"&date_to={to.isoformat()}" if to else "")
    return w.client.get(f"/api/patrol-sheets{q}")


def _cell(body, klass, ordinal, day=WED):
    """某班某節的格子(跨上午/下午與分頁找)。"""
    for page in body["pages"]:
        if page["date"] != day.isoformat() or klass not in page["classes"]:
            continue
        for row in page["rows"]:
            if row["ordinal"] == ordinal:
                return row["cells"][page["classes"].index(klass)]
    raise AssertionError(f"找不到 {klass} 第 {ordinal} 節")


def test_sheet_follows_the_paper_form(w):
    """一天分上午(1~4 節)、下午(5~7 節)兩張;早自習與午休不列;班級為欄。"""
    w.teacher("王師", ["國文"])
    w.place("王師", "國文", "701", 0)
    w.publish()

    r = _sheets(w)
    assert r.status_code == 200, r.json()
    body = r.json()
    assert "115學年第一學期" in body["title"]
    assert "授課情形" in body["legend"]
    assert [(p["first_ordinal"], p["last_ordinal"]) for p in body["pages"]] == [(1, 4), (5, 7)]
    morning = body["pages"][0]
    assert [row["name"] for row in morning["rows"]] == ["第一節", "第二節", "第三節", "第四節"]
    assert morning["table_name"] == ""          # 全校一套節次表:不必寫是哪一套
    assert _cell(body, "701", 1) == {"subject": "國文", "teacher": "王師", "room": "", "note": ""}
    assert _cell(body, "701", 2) == {"subject": "", "teacher": "", "room": "", "note": ""}


def test_substitute_teacher_replaces_the_absent_one(w):
    """代課:那一天印代課老師並註記;其他日子(下週三)仍是原本的老師。"""
    w.teacher("王師", ["國文"])
    w.teacher("陳師", ["國文"])
    w.place("王師", "國文", "701", 0)
    w.publish()
    affected = w.leave("王師")
    assert _cell(_sheets(w).json(), "701", 1)["note"] == "請假待處理"

    status, _ = w.assign(affected[0]["id"], type="substitute",
                         handler_teacher_id=w.teachers["陳師"])
    assert status == 200
    cell = _cell(_sheets(w).json(), "701", 1)
    assert (cell["teacher"], cell["note"]) == ("陳師", "代課")

    next_week = WED + timedelta(days=7)
    cell = _cell(_sheets(w, next_week).json(), "701", 1, next_week)
    assert (cell["teacher"], cell["note"]) == ("王師", "")


def test_self_study_has_no_teacher(w):
    w.teacher("王師", ["國文"])
    w.place("王師", "國文", "701", 0)
    w.publish()
    affected = w.leave("王師")
    status, body = w.assign(affected[0]["id"], type="self_study")
    assert status == 200, body
    cell = _cell(_sheets(w).json(), "701", 1)
    assert (cell["teacher"], cell["note"]) == ("", "自習")


def test_swap_shows_on_both_days(w):
    """調課:請假那節印對調的老師,補課那節印回來補課的老師;兩格都註記調課。"""
    w.teacher("王師", ["國文"])
    w.teacher("陳師", ["數學"])
    w.place("王師", "國文", "701", 0)                 # 週三第一節
    w.place("陳師", "數學", "701", 1, weekday=4)      # 週四第二節
    w.publish()
    affected = w.leave("王師")
    thu = WED + timedelta(days=1)
    _swap_first(w, affected[0]["id"], lambda o: o["date"] == thu.isoformat())

    body = _sheets(w, WED, thu).json()
    leave_cell = _cell(body, "701", 1, WED)
    makeup_cell = _cell(body, "701", 2, thu)
    # 同班互調:科目跟著老師走
    assert (leave_cell["subject"], leave_cell["teacher"], leave_cell["note"]) == (
        "數學", "陳師", "調課")
    assert (makeup_cell["subject"], makeup_cell["teacher"], makeup_cell["note"]) == (
        "國文", "王師", "調課")


def _group(w, name, classes):
    return w.client.post(f"/api/scheduling-units{w.q}", json={
        "name": name, "class_ids": [w.klass(c) for c in classes]}).json()["id"]


def _group_course(w, unit_id, subject, teacher, room=None):
    return w.client.post(f"/api/assignments{w.q}", json={
        "scheduling_unit_id": unit_id, "subject_id": w.subject(subject), "periods_per_week": 1,
        "teachers": [{"teacher_id": w.teachers[teacher]}], "block_rules": [], "room_id": room,
    }).json()["id"]


def _place_group(w, assignment_id, period_idx):
    """排入群組的任一門課,同群組的其他課會一起排進同一時段(H7)。"""
    r = w.client.post(f"/api/timetables/{w.tt}/entries", json={
        "course_assignment_id": assignment_id, "weekday": 3,
        "period_no": w.wed[period_idx]["period_no"], "span": 1})
    assert r.status_code == 201, r.json()


def test_small_group_lists_every_teacher_in_the_cell(w):
    """分組跑班(英語 A、B 兩組):兩班的格子都要看得到兩位老師,不能只印其中一組。"""
    for t in ("英A師", "英B師"):
        w.teacher(t, ["英語"])
    unit = _group(w, "七年級英語分組", ["701", "702"])
    first = _group_course(w, unit, "英語", "英A師")
    _group_course(w, unit, "英語", "英B師")
    _place_group(w, first, 0)
    w.publish()

    body = _sheets(w).json()
    for klass in ("701", "702"):
        cell = _cell(body, klass, 1)
        assert (cell["subject"], cell["teacher"]) == ("英語", "英A師／英B師")
    assert body["group_lists"] == []


def test_big_group_gets_its_own_list(w):
    """社團:全校拆成很多組,格子塞不下,主表只寫群組名稱,另出一張分組巡堂單。"""
    clubs = ["羽球社", "桌球社", "棒球社", "漫畫社"]
    room_id = w.client.post(f"/api/rooms{w.q}", json={"name": "體育館"}).json()["id"]
    unit = _group(w, "社團", ["701", "702"])
    ids = []
    for i, club in enumerate(clubs):
        w.teacher(f"師{i}", [club])
        ids.append(_group_course(w, unit, club, f"師{i}", room=room_id if i == 0 else None))
    _place_group(w, ids[0], 5)
    w.publish()

    body = _sheets(w).json()
    cell = _cell(body, "701", 6)
    assert (cell["subject"], cell["teacher"], cell["note"]) == ("社團", "共 4 組", "見分組巡堂單")
    assert len(body["group_lists"]) == 1          # 兩個班共用一張,不重複
    listing = body["group_lists"][0]
    assert (listing["group_name"], listing["period_name"]) == ("社團", "第六節")
    assert [(i["no"], i["subject"], i["teacher"]) for i in listing["items"]] == [
        (n, club, f"師{n - 1}") for n, club in enumerate(clubs, start=1)]
    assert listing["items"][0]["room"] == "體育館"


def test_group_list_order_follows_the_assignments_not_the_entries(w):
    """分組單的編號照配課建立的順序,與格位寫入資料庫的順序無關。

    群組的格位是整批寫入的,順序取決於資料庫怎麼回傳配課(PostgreSQL 不保證);
    這裡刻意倒著寫入,重現 CI 上「第一列變成棒球社」的情況。
    """
    from app.models.timetable import ScheduleEntry

    clubs = ["羽球社", "桌球社", "棒球社", "漫畫社"]
    unit = _group(w, "社團", ["701", "702"])
    ids = []
    for i, club in enumerate(clubs):
        w.teacher(f"師{i}", [club])
        ids.append(_group_course(w, unit, club, f"師{i}"))
    for assignment_id in reversed(ids):
        w.db.add(ScheduleEntry(
            timetable_id=w.tt, course_assignment_id=assignment_id, weekday=3,
            period_no=w.wed[5]["period_no"], span=1))
    w.db.commit()
    w.publish()

    listing = _sheets(w).json()["group_lists"][0]
    assert [i["subject"] for i in listing["items"]] == clubs


def test_many_classes_are_split_across_pages(w):
    """班級超過一頁的欄數時自動分頁,每頁上午/下午各一張。"""
    w.teacher("王師", ["國文"])
    w.place("王師", "國文", "701", 0)
    for n in range(13):
        w.klass(f"8{n:02d}")
    w.publish()

    pages = _sheets(w).json()["pages"]
    morning = [p for p in pages if p["first_ordinal"] == 1]
    assert [len(p["classes"]) for p in morning] == [12, 3]     # 701、900(佔位班)+ 13 班
    assert len(pages) == 4


def test_days_without_lessons_are_skipped_and_errors_are_explained(w):
    w.teacher("王師", ["國文"])
    w.place("王師", "國文", "701", 0)

    assert _sheets(w).status_code == 404                       # 尚未發布
    w.publish()
    sat = WED + timedelta(days=3)
    r = _sheets(w, sat, sat + timedelta(days=1))               # 週末沒有課
    assert r.status_code == 404 and "沒有任何課" in r.json()["detail"]
    r = _sheets(w, WED, WED + timedelta(days=20))
    assert r.status_code == 404 and "最多" in r.json()["detail"]

    week = _sheets(w, WED - timedelta(days=2), WED + timedelta(days=4)).json()
    assert {p["date"] for p in week["pages"]} == {WED.isoformat()}   # 只有週三排了課


def test_legend_can_be_changed_by_admin(env):
    client, db = env
    make_user(db, "a", PW, roles=[Role.admin])
    client.post("/api/auth/login", json={"username": "a", "password": PW})
    assert "授課情形" in client.get("/api/settings/patrol").json()["legend"]
    r = client.put("/api/settings/patrol", json={"legend": "優、良、可"})
    assert r.json()["legend"] == "優、良、可"
    assert "授課情形" in client.put("/api/settings/patrol", json={"legend": ""}).json()["legend"]


# ── v1.2.14:匯出 Excel ──────────────────────────────────────
def _xlsx(w, day=WED, to=None):
    q = f"{w.q}&date_from={day.isoformat()}" + (f"&date_to={to.isoformat()}" if to else "")
    r = w.client.get(f"/api/patrol-sheets.xlsx{q}")
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    return r, load_workbook(io.BytesIO(r.content))


def _texts(ws) -> list[str]:
    return [str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None]


def test_xlsx_has_a_sheet_per_half_day_with_the_same_content_as_the_print_page(w):
    """上午、下午各一個分頁;代課老師與備註跟列印頁一致,請假的老師不在表上。"""
    w.teacher("王師", ["國文"])
    w.teacher("陳師", ["國文"])
    w.place("王師", "國文", "701", 0)
    w.publish()
    affected = w.leave("王師")
    status, _ = w.assign(affected[0]["id"], type="substitute",
                         handler_teacher_id=w.teachers["陳師"])
    assert status == 200

    r, wb = _xlsx(w)
    day = f"{WED:%m-%d}"
    assert wb.sheetnames == [f"{day} 上午", f"{day} 下午"]
    assert f"UTF-8''%E5%B7%A1%E5%A0%82%E8%A1%A8_{WED.isoformat()}.xlsx" in (
        r.headers["content-disposition"])

    ws = wb[f"{day} 上午"]
    texts = _texts(ws)
    assert "巡堂紀錄" in texts[0] and "第 1~4 節" in texts[0]
    assert f"{WED.year - 1911}年{WED.month}月{WED.day}日" in texts[0]
    header = [c.value for c in ws[2]]
    assert header[0] == "節次" and "701" in header and "巡堂" in header[-1]
    col = header.index("701") + 1
    # 第一節那一塊:科目/教師、教室、授課情形、學生學習、備註
    assert [ws.cell(row=r_, column=2).value for r_ in range(3, 8)] == [
        "科目\n教師", "教室", "授課情形", "學生學習", "備註"]
    assert ws.cell(row=3, column=col).value == "國文\n陳師"
    assert ws.cell(row=7, column=col).value == "代課"
    assert ws.cell(row=5, column=col).value is None        # 授課情形留白給人填
    assert not any("王師" in t for t in texts)
    assert any("授課情形(填寫代碼)" in t for t in texts)     # 記錄說明
    assert ws.page_setup.orientation == "landscape"         # 下載後不改也能直接印


def test_xlsx_puts_each_group_list_on_its_own_sheet(w):
    clubs = ["羽球社", "桌球社", "棒球社", "漫畫社"]
    unit = _group(w, "社團", ["701", "702"])
    ids = []
    for i, club in enumerate(clubs):
        w.teacher(f"師{i}", [club])
        ids.append(_group_course(w, unit, club, f"師{i}"))
    _place_group(w, ids[0], 5)
    w.publish()

    _, wb = _xlsx(w)
    day = f"{WED:%m-%d}"
    assert wb.sheetnames == [f"{day} 上午", f"{day} 下午", f"{day} 社團 第六節"]
    ws = wb[f"{day} 社團 第六節"]
    assert [c.value for c in ws[3]] == ["編號", "名稱", "教師", "地點", "上課狀況", "學生表現"]
    assert [(ws.cell(row=4 + i, column=1).value, ws.cell(row=4 + i, column=2).value)
            for i in range(4)] == list(enumerate(clubs, start=1))
    # 主表那一格只寫群組與組數
    assert any("社團\n共 4 組" == t for t in _texts(wb[f"{day} 下午"]))


def test_xlsx_week_export_and_errors(w):
    w.teacher("王師", ["國文"])
    w.place("王師", "國文", "701", 0)
    assert w.client.get(f"/api/patrol-sheets.xlsx{w.q}&date_from={WED}").status_code == 404
    w.publish()

    mon, sun = WED - timedelta(days=2), WED + timedelta(days=4)
    r, wb = _xlsx(w, mon, sun)
    assert len(wb.sheetnames) == 2                          # 只有週三排了課
    assert f"_{mon.isoformat()}_{sun:%m-%d}.xlsx" in r.headers["content-disposition"]
